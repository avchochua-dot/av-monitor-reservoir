#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""A Vuong DAILY07 V3.9.
Three separate data windows; Excel is populated automatically, not left blank.
Existing SMTP dispatch keys and tables are preserved.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import math
import os
import re
import smtplib
import ssl
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from collections import defaultdict
from copy import copy
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from email.message import EmailMessage
from pathlib import Path
from urllib.error import HTTPError
from urllib.request import Request, urlopen

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel

sys.path.insert(0,str(Path(__file__).resolve().parent))
import generate_daily07_report as daily

BASE=Path(__file__).resolve().parents[1]
TEMPLATE=BASE/"templates"/"A_VUONG_BM_BC_SO_LIEU_HANG_NGAY.xlsx"
DISPATCH_TABLE="av_daily07_bundle_dispatch"
REPORTS={
    "KTTV_31H":("REPORT_KTTV31_TO","pdf"),
    "RESERVOIR_12H":("REPORT_RESERVOIR_TO","pdf"),
    "HYDRO_EXCEL":("REPORT_HYDRO_EXCEL_TO","xlsx"),
}
FLOW_FIELDS=("inflow","turbine_flow","spillway_flow")


def load_data(d, fixture=None):
    raw=json.loads(Path(fixture).read_text(encoding="utf-8")) if fixture else daily.fetch_rows(d)
    if isinstance(raw,dict):
        raw=raw["rows"]
    return raw,daily.validate(raw,d)


def parse_history_date(value):
    if isinstance(value,datetime):
        return value.date()
    if isinstance(value,date):
        return value
    if isinstance(value,str):
        for pattern in ("%d/%m/%Y","%Y-%m-%d","%d-%m-%Y"):
            try:
                return datetime.strptime(value.strip(),pattern).date()
            except ValueError:
                pass
    return None


def _valid_averages(values):
    try:
        numbers=tuple(daily.number(v,"historical mean") for v in values)
    except ValueError:
        return None
    return numbers if all(v>=0 for v in numbers) else None


def read_history(book,cached_book,report_date,source_name):
    ws=book["SolieuQtbngay"]; cached=cached_book["SolieuQtbngay"]
    history={}; audit=[]; occurrences=defaultdict(list)
    for r in range(6,ws.max_row+1):
        d=parse_history_date(ws.cell(r,2).value)
        if d is None or d>=report_date:
            continue
        values=[]
        for c in (3,4,5):
            cell=ws.cell(r,c)
            values.append(cached.cell(r,c).value if cell.data_type=="f" else cell.value)
        occurrences[d].append((r,tuple(values)))
    for d,records in occurrences.items():
        # User-approved dedup rule: keep last row, do not sum/average duplicates.
        retained_row,retained_values=records[-1]
        valid=_valid_averages(retained_values)
        if valid is not None:
            history[d]={"values":valid,"source":f"{source_name}:SolieuQtbngay!B{retained_row}:E{retained_row}"}
        for row,values in records[:-1]:
            audit.append([d, row, retained_row,*values,values!=retained_values])
    return history,audit


def dates_between(a,b):
    return [a+timedelta(days=i) for i in range(max(0,(b-a).days+1))]


def contiguous(groups):
    if not groups:
        return []
    ordered=sorted(set(groups)); result=[]; start=last=ordered[0]
    for d in ordered[1:]:
        if d!=last+timedelta(days=1):
            result.append((start,last)); start=d
        last=d
    result.append((start,last))
    return result


def complete_history(raw,history,report_date,fixture=False,backfill_from=None):
    """Rebuild unseeded days from hourly Supabase data, even on fresh CI runners.
    The shipped historical workbook is an explicitly sourced seed. Every day
    after its watermark is recomputed per run; no in-memory persistence assumed.
    """
    end=report_date-timedelta(days=1)
    week_start=end-timedelta(days=6)
    last=max(history,default=week_start-timedelta(days=1))
    required=set(dates_between(week_start,end))
    required.update(dates_between(last+timedelta(days=1),end))
    if backfill_from is not None:
        if backfill_from>end:
            raise ValueError("--backfill-from is later than data day")
        required.update(dates_between(backfill_from,end))
    targets=sorted({d for d in required if d not in history} | {end})
    # Never fabricate zero flow for a missing day.
    combined=list(raw); selected_raw=daily.latest_rows(combined)
    need_fetch=[]
    for d in targets:
        for t in daily.hydro_times(d):
            row=selected_raw.get(t)
            if row is None or any(row.get(k) is None for k in FLOW_FIELDS):
                need_fetch.append(d);break
    if need_fetch and not fixture:
        for start_day,end_day in contiguous(need_fetch):
            combined.extend(daily.fetch_range(daily.hydro_times(start_day)[0],
                                              daily.hydro_times(end_day)[-1]))
    filled={}
    for d in targets:
        hourly=daily.select_exact(combined,daily.hydro_times(d),FLOW_FIELDS)
        avgs=tuple(daily.mean(hourly,k) for k in FLOW_FIELDS)
        filled[d]=hourly
        history[d]={"values":avgs,"source":"reservoir_hourly_data; 08h D -> 07h D+1"}
    if any(d not in history for d in dates_between(week_start,end)):
        raise ValueError("PENDING_DATA: weekly period does not contain seven complete days")
    return history,filled,combined


def _set_formula(ws,address,formula,value,cache):
    ws[address]=formula
    cache[(ws.title,address)]=value


def _font_name(cell):
    if cell.__class__.__name__!="MergedCell":
        f=copy(cell.font);f.name="Times New Roman";cell.font=f


def _put_cached_values(path,cache):
    """Save cached results for only formulas authored and independently calculated
    by this program. Formulas remain editable Excel formulas; no formula is removed.
    No office installation or proprietary rendering package is needed at runtime.
    """
    ns={"s":"http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r":"http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    main=ns["s"]
    with zipfile.ZipFile(path,"r") as archive:
        contents={info.filename:archive.read(info.filename) for info in archive.infolist()}
    wb=ET.fromstring(contents["xl/workbook.xml"])
    rel=ET.fromstring(contents["xl/_rels/workbook.xml.rels"])
    links={r.attrib["Id"]:r.attrib["Target"] for r in rel}
    applied=set()
    for sheet in wb.find("s:sheets",ns):
        name=sheet.attrib["name"]
        target=links[sheet.attrib["{"+ns["r"]+"}id"]]
        xmlpath=target.lstrip("/") if target.startswith("/") else "xl/"+target
        tree=ET.fromstring(contents[xmlpath])
        for cell in tree.findall(".//s:sheetData/s:row/s:c",ns):
            address=cell.attrib["r"];key=(name,address)
            if cell.find("s:f",ns) is None:
                continue
            if key not in cache:
                raise RuntimeError("UNCACHED_FORMULA: "+name+"!"+address)
            value=cache[key];applied.add(key)
            old=cell.find("s:v",ns)
            if old is None:
                old=ET.SubElement(cell,"{"+main+"}v")
            if isinstance(value,str):
                cell.set("t","str");old.text=value
            elif isinstance(value,bool):
                cell.set("t","b");old.text="1" if value else "0"
            elif value is None:
                cell.set("t","str");old.text=""
            else:
                v=float(value)
                if not math.isfinite(v):
                    raise ValueError("Cannot cache non-finite formula result")
                cell.attrib.pop("t",None)
                old.text=format(v,".15g")
        contents[xmlpath]=ET.tostring(tree,encoding="utf-8",xml_declaration=True)
    if set(cache)!=applied:
        raise RuntimeError("Formula cache mismatch: "+str(set(cache)-applied))
    temp=Path(str(path)+".tmp")
    with zipfile.ZipFile(temp,"w",zipfile.ZIP_DEFLATED) as out:
        for name,data in contents.items():
            out.writestr(name,data)
    temp.replace(path)


def update_excel(raw,rows,report_date,target,template=TEMPLATE,fixture=False,backfill_from=None):
    if not Path(template).is_file():
        raise FileNotFoundError(f"Missing Excel template: {template}")
    w=openpyxl.load_workbook(template)
    cached=openpyxl.load_workbook(template,data_only=True)
    required_sheets={"Solieuhangngay_gio","SolieuQtbngay","Tuan"}
    if not required_sheets.issubset(w.sheetnames):
        raise ValueError("The workbook is not the approved A Vuong template")
    history,duplicates=read_history(w,cached,report_date,Path(template).name)
    original_dates=set(history)
    history,filled,raw_all=complete_history(raw,history,report_date,fixture,backfill_from)
    prev=report_date-timedelta(days=1)
    # Window independent of the daily calendar statistics in the reservoir PDF.
    hourly=daily.select_exact(raw_all,daily.hydro_times(prev),daily.FIELDS)
    cache={}
    sh=w["Solieuhangngay_gio"]
    # Exactly the two groups in the supplied original: 08..24 and 01..07.
    for merged in list(sh.merged_cells.ranges):
        if merged.min_col==2 and merged.max_col==2 and merged.max_row>=6 and merged.min_row<=29:
            sh.unmerge_cells(str(merged))
    for r in range(6,30):
        sh.cell(r,2).value=None
    for group,value in [("B6:B22",prev),("B23:B29",report_date)]:
        sh.merge_cells(group)
        c=sh[group.split(":")[0]]
        c.value=datetime.combine(value,datetime.min.time())
        c.number_format="dd/mm/yyyy"
        c.alignment=Alignment(horizontal="center",vertical="center")
    for i,(t,r) in enumerate(hourly,6):
        sh.cell(i,3,24 if t.hour==0 else t.hour)
        sh.cell(i,3).number_format="0"
        for col,key in [(4,"water_level"),(5,"inflow"),(6,"turbine_flow"),(7,"spillway_flow")]:
            cell=sh.cell(i,col)
            cell.value=float(daily.number(r[key],key))
            cell.number_format="0.00"
            cell.comment=Comment(f"Nguồn: reservoir_hourly_data.{key}\nGiờ nhãn Việt Nam: {t.isoformat(' ')}",
                                 "A Vuong data pipeline")
    sh["B1"]=report_date.year
    _set_formula(sh,"B3",'="Số liệu thủy văn theo giờ từng ngày năm "&B1',
                 f"Số liệu thủy văn theo giờ từng ngày năm {report_date.year}",cache)

    # Hidden raw-input sheet provides traceable formulas for all backfilled dates.
    for name in ["_NguonGio","_DoiSoatNguon"]:
        if name in w.sheetnames:
            del w[name]
    source=w.create_sheet("_NguonGio")
    source.append(["Giờ nhãn VN","Mực nước hồ","Mực nước hạ lưu",
                   "Q về","Q máy","Q xả tràn","Mưa giờ"])
    source_map={}
    source_rows={}
    for period in filled.values():
        for t,r in period:
            source_rows[t]=r
    for t in sorted(source_rows):
        r=source_rows[t]
        source.append([t]+[float(r[k]) if r.get(k) is not None else None for k in daily.FIELDS])
        source_map[t]=source.max_row
        source.cell(source.max_row,1).number_format="dd/mm/yyyy hh:mm"
    source.sheet_state="hidden"

    avgsh=w["SolieuQtbngay"]
    avgsh["B3"]=f"Số liệu thủy văn theo ngày năm {report_date.year}"
    prototype={col:copy(avgsh.cell(6,col)._style) for col in range(2,6)}
    # Rebuild the generated sheet's dated range; never edit Supabase history.
    old_max=avgsh.max_row
    for row in avgsh.iter_rows(min_row=6,max_row=old_max,min_col=2,max_col=5):
        for cell in row:
            cell.value=None;cell.comment=None
    index={}
    for rn,d in enumerate(sorted(history),6):
        index[d]=rn
        avgsh.row_dimensions[rn].height=19.5
        for col in range(2,6):
            avgsh.cell(rn,col)._style=copy(prototype[col])
        avgsh.cell(rn,2,datetime.combine(d,datetime.min.time()))
        avgsh.cell(rn,2).number_format="dd/mm/yyyy"
        vals=history[d]["values"]
        for c,k,v,rawcol in zip((3,4,5),FLOW_FIELDS,vals,("D","E","F")):
            coord=f"{get_column_letter(c)}{rn}"
            if d in filled:
                times=daily.hydro_times(d)
                first,last=source_map[times[0]],source_map[times[-1]]
                if last-first!=23:
                    raise AssertionError("Backfilled source period must contain 24 consecutive rows")
                formula=f"=SUM('_NguonGio'!{rawcol}{first}:{rawcol}{last})/24"
                _set_formula(avgsh,coord,formula,float(v),cache)
            else:
                avgsh[coord]=float(v)
            avgsh[coord].number_format="0.00"
            avgsh[coord].comment=Comment(history[d]["source"],"A Vuong data pipeline")
    lastrow=max(index.values())
    # Retain 3 visible report sheets; archive removed historical duplicate rows.
    audit=w["Duplicate_Audit"] if "Duplicate_Audit" in w else w.create_sheet("Duplicate_Audit")
    if audit.max_row==1 and audit["A1"].value is None:
        audit.append(["Ngày","Dòng bị loại trong mẫu","Dòng được giữ","Q về","Q máy",
                      "Q xả tràn","Khác giá trị với dòng giữ"])
    for rec in duplicates:
        audit.append(rec)
    audit.sheet_state="hidden"
    provenance=w.create_sheet("_DoiSoatNguon")
    provenance.append(["Ngày","Nguồn số liệu","Khoảng tính trung bình","Trạng thái"])
    for d in sorted(history):
        provenance.append([d,history[d]["source"],"08h D -> 07h D+1; 24 mốc",
                           "TÍNH TỪ GIỜ" if d in filled else "LỊCH SỬ TRONG FILE NGUỒN"])
    provenance.sheet_state="hidden"

    week=w["Tuan"];week["A2"]=report_date.year
    week["B1"]=datetime(report_date.year,1,1)
    week["B1"].number_format="dd/mm/yyyy"
    week_days=daily_dates=dates_between(prev-timedelta(days=6),prev)
    values_week=[history[d]["values"] for d in week_days]
    _set_formula(week,"A3","=ISOWEEKNUM(C12)",prev.isocalendar().week,cache)
    _set_formula(week,"F1","=WEEKDAY(C6,1)",(week_days[0].weekday()+1)%7+1,cache)
    _set_formula(week,"B3",'="Số liệu thủy văn tuần thứ "&A3&" năm "&A2',
                 f"Số liệu thủy văn tuần thứ {prev.isocalendar().week} năm {report_date.year}",cache)
    weekdays=["Thứ hai","Thứ ba","Thứ tư","Thứ năm","Thứ sáu","Thứ bảy","Chủ nhật"]
    for rn,d in enumerate(week_days,6):
        week.cell(rn,3,datetime.combine(d,datetime.min.time()))
        week.cell(rn,3).number_format="dd/mm/yyyy"
        _set_formula(week,f"B{rn}",
          f'=CHOOSE(WEEKDAY(C{rn},2),"Thứ hai","Thứ ba","Thứ tư","Thứ năm","Thứ sáu","Thứ bảy","Chủ nhật")',
          weekdays[d.weekday()],cache)
        for c,idx,v in zip((4,5,6),(2,3,4),history[d]["values"]):
            formula=(f'=IF(COUNTIF(SolieuQtbngay!$B$6:$B${lastrow},$C{rn})=1,'
                     f'VLOOKUP($C{rn},SolieuQtbngay!$B$6:$E${lastrow},{idx},FALSE),"THIEU_DU_LIEU")')
            _set_formula(week,f"{get_column_letter(c)}{rn}",formula,float(v),cache)
            week.cell(rn,c).number_format="0.00"
    _set_formula(week,"B13",'="Đặc trưng"&CHAR(10)&"Tuần "&A3',
                 f"Đặc trưng\nTuần {prev.isocalendar().week}",cache)
    for c,ix in zip("DEF",range(3)):
        vals=[v[ix] for v in values_week]
        summary=[max(vals),sum(vals,Decimal(0))/7,min(vals)]
        for rn,func,result in zip((13,14,15),("MAX","AVERAGE","MIN"),summary):
            _set_formula(week,f"{c}{rn}",
                         f'=IF(COUNT({c}6:{c}12)=7,{func}({c}6:{c}12),"THIEU_DU_LIEU")',
                         float(result),cache)
            week[f"{c}{rn}"].number_format="0.00"
    # Style of the original template is retained; only font family is standardized.
    for ws in w:
        for row in ws:
            for cell in row:
                if cell.value is not None or cell.has_style:
                    _font_name(cell)
        ws.sheet_view.showGridLines=False
    # Preserve template widths, fills, sizes, bold and italic; set coherent print areas.
    sh.print_area="B2:G29";sh.print_title_rows="2:5"
    avgsh.print_area=f"B2:E{lastrow}";avgsh.print_title_rows="2:5"
    week.print_area="B2:F15"
    for ws in (sh,avgsh,week):
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation="portrait";ws.page_setup.paperSize=ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.page_margins.left=0.5;ws.page_margins.right=0.5
        ws.page_margins.top=0.5;ws.page_margins.bottom=0.5
    # Exact cached values are written after save; Excel will still recalculate edits.
    w.calculation=openpyxl.workbook.properties.CalcProperties(
        calcMode="auto",fullCalcOnLoad=True,forceFullCalc=True)
    w.save(target)
    _put_cached_values(target,cache)
    # Verify that all formulas have valid cached results in the emitted file.
    check=openpyxl.load_workbook(target,data_only=True)
    for (sheet,cell),value in cache.items():
        result=check[sheet][cell].value
        if result is None:
            raise RuntimeError(f"UNCACHED_RESULT: {sheet}!{cell}")
        if isinstance(value,(int,float,Decimal)) and abs(float(result)-float(value))>1e-8:
            raise RuntimeError(f"CACHE_MISMATCH: {sheet}!{cell}")
        if isinstance(result,str) and result.startswith(("#REF!","#DIV/0!","#VALUE!","#NAME?","#N/A")):
            raise RuntimeError(f"FORMULA_ERROR: {sheet}!{cell}")
    summary={
        "hydro_window":[hourly[0][0].isoformat(),hourly[-1][0].isoformat()],
        "hydro_hour_count":len(hourly),
        "daily_hydro_means":dict(zip(FLOW_FIELDS,map(float,history[prev]["values"]))),
        "history_days":len(history),
        "history_added_days":[d.isoformat() for d in sorted(set(history)-original_dates)],
        "history_recalculated_days":[d.isoformat() for d in sorted(filled)],
        "historical_duplicates_archived":len(duplicates),
        "week_dates":[d.isoformat() for d in week_days],
        "week_complete_days":len(values_week),
        "formulas_with_cached_results":len(cache),
        "source_template":Path(template).name,
    }
    print("EXCEL_DATA_READY",json.dumps(summary,ensure_ascii=False))
    return summary


def recipients_for(report):
    env,_=REPORTS[report]
    raw=os.getenv(env,"")
    return sorted({x.strip().lower() for x in raw.replace(";",",").split(",") if x.strip()})

def dispatch_email_once(report, report_date, file, recipients):
    """Atomic DB INSERT claim, no automatic SMTP retry on ambiguous failure."""
    if not recipients:
        print(f"EMAIL_NOT_CONFIGURED {report}; skipped")
        return
    for k in ("SMTP_HOST","SMTP_USER","SMTP_PASSWORD","SMTP_FROM"):
        if not os.getenv(k): raise ValueError("Missing " + k)
    # One deterministic claim for each report type and date.
    # Recipient changes must NOT trigger a second automatic send.
    report_key = f"test:daily07:{report}:{report_date.isoformat()}"
    claim={
        "report_key": report_key,
        "report_type": report,
        "report_date": report_date.isoformat(),
        "status": "SENDING",
        "recipients": ",".join(recipients),
        "file_sha256": hashlib.sha256(file.read_bytes()).hexdigest()
    }
    # The database must enforce UNIQUE(report_type, report_date).
    result=daily.rest_request("POST",DISPATCH_TABLE,claim,"return=representation")
    if isinstance(result,dict) and result.get("_conflict"):
        print("EMAIL_SKIPPED_ALREADY_CLAIMED",report,report_date)
        return
    msg=EmailMessage()
    titles = {
        "KTTV_31H": "Báo cáo số liệu vận hành hồ thủy điện A Vương",
        "RESERVOIR_12H": "Báo cáo số liệu vận hành Nhà máy thủy điện A Vương",
        "HYDRO_EXCEL": "Bảng số liệu thủy văn hồ thủy điện A Vương",
    }
    msg["Subject"] = f"[KIỂM TRA NỘI BỘ] {titles[report]} - 07h00 ngày {report_date:%d/%m/%Y}"
    msg["From"]=os.environ["SMTP_FROM"]
    msg["To"]=", ".join(recipients)
    msg.set_content(
        "Kính gửi anh/chị,\n\n"
        "Đính kèm báo cáo số liệu vận hành A Vương lúc 07h00 ngày "
        f"{report_date:%d/%m/%Y}, theo biểu mẫu vận hành.\n"
        "Email này chỉ gửi tới danh sách kiểm tra nội bộ. "
        "Văn bản chưa có chữ ký của người có thẩm quyền.\n"
        "Trân trọng."
    )
    ispdf=file.suffix==".pdf"
    msg.add_attachment(file.read_bytes(),
      maintype="application",
      subtype="pdf" if ispdf else "vnd.openxmlformats-officedocument.spreadsheetml.sheet",
      filename=file.name)
    try:
        security=os.getenv("SMTP_SECURITY","ssl").lower()
        port=int(os.getenv("SMTP_PORT","465"))
        ctx=ssl.create_default_context()
        if security=="ssl":
            conn=smtplib.SMTP_SSL(os.environ["SMTP_HOST"],port,timeout=35,context=ctx)
        elif security=="starttls":
            conn=smtplib.SMTP(os.environ["SMTP_HOST"],port,timeout=35)
        else: raise ValueError("SMTP_SECURITY must be ssl or starttls")
        with conn as s:
            if security=="starttls": s.starttls(context=ctx)
            s.login(os.environ["SMTP_USER"],os.environ["SMTP_PASSWORD"])
            refusal=s.send_message(msg)
            if refusal: raise RuntimeError("Some recipients refused by SMTP")
    except Exception as e:
        # SMTP may have accepted before a transport failure; manual review only.
        try:
            daily.rest_request("PATCH",
                f"{DISPATCH_TABLE}?report_key=eq.{report_key}",
                {"status":"NEEDS_REVIEW","error_text":str(e)[:500]})
        finally: raise
    daily.rest_request("PATCH",
      f"{DISPATCH_TABLE}?report_key=eq.{report_key}",
      {"status":"SMTP_ACCEPTED","sent_at":datetime.now(timezone.utc).isoformat()})
    print(f"EMAIL_SMTP_ACCEPTED {report} recipients={len(recipients)}")





def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--date",required=True)
    ap.add_argument("--fixture",help="Offline hourly JSON; SMTP disabled")
    ap.add_argument("--template",default=str(TEMPLATE))
    ap.add_argument("--backfill-from",help="Optional date to rebuild missing hydro-history days")
    ap.add_argument("--output-dir",default="reports/avuong_kttv/reports/output")
    ap.add_argument("--send-test-email",action="store_true")
    args=ap.parse_args()
    report_date=daily.parse_day(args.date)
    if args.fixture and args.send_test_email:
        raise ValueError("Do not send email with fixture data")
    if not args.fixture and datetime.combine(report_date,datetime.min.time()).replace(
            hour=7,tzinfo=daily.VN)>datetime.now(daily.VN):
        raise ValueError("Report issue is in the future")
    raw,rows=load_data(report_date,args.fixture)
    out=Path(args.output_dir);out.mkdir(parents=True,exist_ok=True)
    files={
        "KTTV_31H":out/f"07h00_{report_date:%d-%m-%Y}_So_lieu_KTTV_tong_hop.pdf",
        "RESERVOIR_12H":out/f"07h00_{report_date:%d-%m-%Y}_Bao_cao_van_hanh_ho.pdf",
        "HYDRO_EXCEL":out/f"07h00_{report_date:%d-%m-%Y}_Bang_so_lieu_thuy_van.xlsx",
    }
    # Validate/populate workbook before any SMTP side effect.
    excel=update_excel(raw,rows,report_date,files["HYDRO_EXCEL"],
                       template=Path(args.template),fixture=bool(args.fixture),
                       backfill_from=daily.parse_day(args.backfill_from) if args.backfill_from else None)
    daily.build_pdf(rows,report_date,files["KTTV_31H"])
    hours=daily.render_reservoir_pdf(rows,report_date,files["RESERVOIR_12H"])
    manifest={
        "version":daily.VERSION,"report_date":report_date.isoformat(),"status":"BUNDLE_READY",
        "pdf_day_window":[rows[0][0].isoformat(),rows[23][0].isoformat()],
        "pdf_day_means":{k:float(daily.mean(rows[:24],k)) for k in FLOW_FIELDS},
        "generation_hours":hours,"excel":excel,
        "files":{k:{"name":p.name,"sha256":hashlib.sha256(p.read_bytes()).hexdigest()}
                 for k,p in files.items()},
    }
    (out/f"DAILY07_{report_date:%Y%m%d}_AUDIT.json").write_text(
        json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf-8")
    print(f"BUNDLE_READY report_date={report_date} rows=31 generation_hours={hours}")
    for label,file in files.items():
        print("FILE_READY",label,file)
    if args.send_test_email:
        for label,file in files.items():
            dispatch_email_once(label,report_date,file,recipients_for(label))
    else:
        print("PREVIEW_ONLY_NO_EMAIL")


if __name__=="__main__":
    main()
